from pathlib import Path
import json, math
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from build_report_html import fill_stats, build_report_from_markdown, candidate_table, kpis_from_stats

ROOT=Path(__file__).resolve().parents[1];D=ROOT/'data';STUDY='Liver-Lipid-GPT'
def read(name):return pd.read_csv(D/(name+'.tsv'),sep='\t')
stats=json.loads((D/'stats.json').read_text());s=read('assay_summary');e=read('enrichment');ranked=json.loads((D/'ranked_results.json').read_text())
def mdtable(df):
    def val(v):
        if isinstance(v,float):return f'{v:.3g}'
        return str(v).replace('|',' / ').replace('\n',' ')
    return '| '+' | '.join(df.columns)+' |\n|'+ '|'.join(['---']*len(df.columns))+'|\n'+'\n'.join('| '+' | '.join(val(v) for v in row)+' |' for row in df.itertuples(index=False,name=None))
beta=e[(e.study=='OSD-25')&(e.rule=='sensitivity')&(e.direction=='up')&(e.background=='all annotated')&(e.label=='fatty acid beta-oxidation')].iloc[0]
fatty=e[(e.study=='OSD-25')&(e.rule=='strict')&(e.direction=='all')&(e.background=='retained conditional')&(e.label=='non-alcoholic fatty liver disease')].iloc[0]
for prefix,row in [('beta',beta),('fatty',fatty)]:
    for k in ['k','K','n','N','fold','fdr']:stats[prefix+'_'+k]=int(row[k]) if k in ['k','K','n','N'] else float(row[k])
cov=s[['study','retained','fdr05','strict','sensitivity']].copy();cov.insert(1,'Cohort',['STS-135, 13 days','RR1-CASIS, 21 days','RR3, 42 days']);cov.columns=['Dataset','Cohort','Retained','Adjusted p ≤ 0.05','Strict twofold','Sensitivity 1.2-fold']
stats['coverage_table']=mdtable(cov)
tr=e[(e.study=='OSD-25')&(e.rule=='strict')&(e.direction=='all')&(e.background=='retained conditional')&(e.family=='DIG lipid traits')][['label','k','K','n','N','expected','fold','fdr']].copy();stats['trait_table']=mdtable(tr)
rs=pd.DataFrame(ranked).head(8)[['mouse_genes','human_entrez','study','log2fc','adj_p','tier','sources_n']];stats['ranked_slice']=mdtable(rs)
(D/'stats.json').write_text(json.dumps(stats,indent=2))
md=ROOT/(STUDY+'_report.md');md.write_text(fill_stats((D/'report_template.md').read_text(),stats))
columns=[('rank','Rank'),('mouse_genes','Mouse gene(s)'),('human_entrez','Human Entrez'),('study','Study'),('log2fc','log₂FC'),('adj_p','Adjusted p'),('tier','Tier'),('sources_n','Sources (n)'),('lipid_terms','Lipid/circadian annotation'),('traits','Trait links'),('ambiguity','Mapping')]
table=candidate_table(ranked,columns,numeric_keys=['rank','log2fc','adj_p','sources_n'],extra_filters=[('study','Study'),('direction','Direction'),('tier','Tier'),('ambiguity','Mapping')],sources_col=('sources_n','sources'),default_sort='rank')
build_report_from_markdown(md,ROOT/(STUDY+'_report.html'),stats=stats,kpis=kpis_from_stats(stats,[('assays','Clean paper-linked contrasts'),('retained_total','Retained gene–assay records'),('strict_total','Strict mouse gene–assay results'),('human_rows','Ranked human-ortholog rows')]),table=table)

wb=Workbook();wb.remove(wb.active)
def cellval(v):
    if isinstance(v,(list,dict)):return json.dumps(v)
    if v is None:return None
    try:
        if pd.isna(v):return None
    except (ValueError,TypeError):pass
    if hasattr(v,'item'):v=v.item()
    return v
def sheet(name,df,widths=None):
    ws=wb.create_sheet(name);ws.append(list(df.columns))
    for row in df.itertuples(index=False,name=None):ws.append([cellval(v) for v in row])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws.row_dimensions[1].height=32
    for c in ws[1]:c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F3864');c.alignment=Alignment(wrap_text=True,vertical='center')
    for j,col in enumerate(df.columns,1):
        width=18
        if col in ['label','lipid_terms','diseases','traits','category','gene','human','g2','term','assay']:width=48
        if col in ['mouse_genes','mouse_symbols','rule','background','family']:width=25
        ws.column_dimensions[get_column_letter(j)].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(c.value,float):c.number_format='0.000E+00' if abs(c.value)<.001 and c.value!=0 else '0.000'
    return ws
rr=pd.DataFrame(ranked);ws=sheet('Ranked Results',rr)
tcol=list(rr.columns).index('tier')+1
for row in ws.iter_rows(min_row=2):
    color='E3F0E8' if row[tcol-1].value=='B' else 'F3F5F7'
    for c in row:c.fill=PatternFill('solid',fgColor=color)
    ws.row_dimensions[row[0].row].height=60
sheet('Assay Inventory',cov)
sheet('Analysis Status',read('analysis_status'))
sheet('Mouse Expression',read('mouse_results'))
sheet('Human Orthologs',read('human_results'))
for family,name in [('GO BP','GO BP Enrichment'),('Reactome','Reactome Enrichment'),('RDKG curated genetic','Genetic Disease Enrichment'),('DIG lipid traits','Lipid Trait Enrichment')]:
    # All tests, including zero overlaps, remain in data/enrichment.tsv.
    # Workbook includes all nonzero-overlap tests, even nonsignificant results.
    sheet(name,e[(e.family==family)&(e.k>0)])
sheet('Lipid Functional Links',read('lipid_annotations'))
for name,title in [('disease_bridge','Disease Links'),('trait_bridge','Trait Links')]:
    j=json.loads((D/(name+'.json')).read_text());sheet(title,pd.DataFrame(j['data'],columns=j['columns']))
methods=[
 ('Scope','Partial thematic reproduction of Beheshti et al. 2019; selected OSD-25, OSD-47, OSD-137 clean contrasts.'),
 ('Original paper','PubMed PMID 31844325; https://doi.org/10.1038/s41598-019-55869-2; full text via Paperclip.'),
 ('Graph selection','Adjusted p ≤ 0.1, per https://github.com/BaranziniLab/spoke_genelab. Omitted genes are unobserved in the filtered graph.'),
 ('Strict rule','Adjusted p ≤ 0.05; absolute log2 fold change ≥ 1.'),
 ('Sensitivity rule','Adjusted p ≤ 0.05; absolute log2 fold change ≥ log2(1.2). Not an exact original-pipeline recreation.'),
 ('Contrasts','Space Flight arm 1, Ground Control arm 2; matched tissue and stripped covariates; no pooling.'),
 ('Orthologs','Per assay max absolute effect with corresponding p; mean-effect sign and mapping ambiguity retained.'),
 ('Functional joins','Mouse→human Entrez in spoke-genelab; Wikidata P351→P594/P354; ProKN encodes→GO/Reactome. Aliases collapsed to human Entrez.'),
 ('Disease joins','Human Entrez normalized to identifiers.org/ncbigene for rdkg. General related_to annotation separated from genetic_association ORA.'),
 ('Trait joins','Human Entrez→digcfdekg geneToTrait; four selected lipid/fatty-liver sets; background all 21,710 trait genes, plus conditional retained universe.'),
 ('ORA','One-sided hypergeometric; BH includes zero-overlap eligible categories. Separate family per assay/rule/direction/background.'),
 ('Workbook rows','Enrichment sheets contain all nonzero-overlap rows; complete test universe, including zero overlaps, is data/enrichment.tsv.'),
 ('Interpretation','Ortholog-inferred, associative and selected-universe results. No causal, clinical, treatment or pathway-activation claim.'),
 ('Ranking','10×lipid/circadian membership + biological source count + min(abs(log2FC),5)/10; p-value tie-break. Tier B requires unambiguous relevant annotation, otherwise C; no A.'),
 ('Reproduction','Run scripts/analyze.py, scripts/figures.py and scripts/build_package.py using Python with pandas, numpy, scipy, matplotlib and openpyxl; data extracts are included.'),
 ('Abbreviations','KG: knowledge graph; GO BP: Gene Ontology biological process; FDR: false discovery rate; BH: Benjamini–Hochberg; ORA: over-representation analysis; OSD: Open Science Data; NAFLD: non-alcoholic fatty liver disease.'),
 ('Model','GPT-6-Astra, provided by the user.'),
 ('Source versions','spoke-genelab v0.0.2; prokn v0.0.5; rdkg v0.0.1; digcfdekg v0.0.1; Wikidata version unavailable in registry.')]
mw=sheet('Methods & Rules',pd.DataFrame(methods,columns=['Item','Rule / source']));mw.column_dimensions['A'].width=27;mw.column_dimensions['B'].width=110
for i in range(2,mw.max_row+1):mw.row_dimensions[i].height=44
out=ROOT/(STUDY+'_results.xlsx');wb.save(out)
check=load_workbook(out,read_only=True,data_only=True)
assert check['Ranked Results'].max_row==len(ranked)+1
assert check['Mouse Expression'].max_row==stats['retained_total']+1
print('Workbook saved and row counts verified:',out)
